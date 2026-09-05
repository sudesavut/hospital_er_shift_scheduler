"""Tests for the Özet (summary) sheet's day/night shift-count breakdown."""

from pathlib import Path

import openpyxl

from nobet_scheduler.doctor import load_doctors
from nobet_scheduler.export import export_schedule
from nobet_scheduler.model import build_and_solve_schedule

FIXTURE = Path(__file__).parent.parent / "data" / "doctors_test_scale.csv"

EXPECTED_HEADERS = [
    "Doktor",
    "Kıdem",
    "Hedef Nöbet Sayısı",
    "Atanan Nöbet Sayısı",
    "Gündüz Nöbet Sayısı",
    "Gece Nöbet Sayısı",
]


def _export_and_read_summary(tmp_path):
    doctors = load_doctors(FIXTURE)
    result = build_and_solve_schedule(doctors, year=2026, month=9, seed=42, max_time_seconds=60)
    output_path = tmp_path / "nobet_listesi.xlsx"
    export_schedule(result, doctors, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Özet"]
    return doctors, ws


def test_summary_sheet_has_day_night_columns_in_expected_order(tmp_path):
    _, ws = _export_and_read_summary(tmp_path)
    header_row = [cell.value for cell in ws[1]]
    assert header_row == EXPECTED_HEADERS


def test_summary_sheet_day_plus_night_equals_assigned_for_every_doctor(tmp_path):
    doctors, ws = _export_and_read_summary(tmp_path)
    doctors_by_name = {d.name: d for d in doctors}

    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, seniority, target, assigned, day_count, night_count = row
        seen.add(name)

        assert day_count + night_count == assigned, f"{name}: gündüz+gece != atanan"
        assert assigned == doctors_by_name[name].shift_count_target, (
            f"{name}: hedef ile atanan uyuşmuyor (hard constraint ihlali)"
        )
        assert day_count >= 0 and night_count >= 0

    assert seen == set(doctors_by_name)


def test_summary_sheet_zero_target_doctor_has_zero_day_and_night(tmp_path):
    doctors, ws = _export_and_read_summary(tmp_path)
    zero_target_names = {d.name for d in doctors if d.shift_count_target == 0}
    assert zero_target_names, "fixture should contain at least one zero-target doctor"

    for row in ws.iter_rows(min_row=2, values_only=True):
        name, _seniority, _target, assigned, day_count, night_count = row
        if name in zero_target_names:
            assert (assigned, day_count, night_count) == (0, 0, 0)

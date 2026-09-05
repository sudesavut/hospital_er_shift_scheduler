"""Export a solved schedule to nobet_listesi.xlsx.

Sheet 1 ("Nöbet Listesi"): rows = dates, columns = gündüz/gece x role.
Sheet 2 ("Özet"): per-doctor total assigned shifts vs. target.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .calendar_utils import Shift
from .doctor import Doctor
from .model import ROLE_COMEZ, ROLE_COMEZ_BASI, ROLE_KAPICI, ROLE_NOBET_BASI, ScheduleResult

ROLE_LABELS = {
    ROLE_NOBET_BASI: "Nöbet Başı",
    ROLE_KAPICI: "Kapıcı",
    ROLE_COMEZ_BASI: "Çömez Başı",
}

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GROUP_HEADER_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")

FOOTER_LINE_1 = "Hazırlayan: S. Savut"
FOOTER_LINE_2 = "Sorularınız için: ssavut@ethz.ch"
FOOTER_FONT = Font(size=8, color="808080")


def _write_footer(ws: Worksheet, last_used_row: int) -> None:
    """Writes the two-line footer a few blank rows below the table, in a
    small/plain font so it reads as a signature, not table data."""
    row = last_used_row + 3
    ws.cell(row=row, column=1, value=FOOTER_LINE_1).font = FOOTER_FONT
    ws.cell(row=row + 1, column=1, value=FOOTER_LINE_2).font = FOOTER_FONT


def export_schedule(result: ScheduleResult, doctors: list[Doctor], output_path: str | Path) -> None:
    wb = Workbook()
    _write_schedule_sheet(wb.active, result)
    _write_summary_sheet(wb.create_sheet("Özet"), result, doctors)
    wb.save(output_path)


def _shift_columns(result: ScheduleResult, is_day_shift: bool) -> list[tuple[str, str]]:
    """Ordered (column_key, header_label) pairs for one shift type (day/night)."""
    shifts = [s for s in result.shifts if s.is_day_shift == is_day_shift]

    max_comez_basi = max(
        (len(result.assignments.get((s, ROLE_COMEZ_BASI), [])) for s in shifts), default=1
    )
    max_comez_basi = max(max_comez_basi, 1)
    max_comez = max((len(result.assignments.get((s, ROLE_COMEZ), [])) for s in shifts), default=0)

    columns = [(f"{ROLE_NOBET_BASI}#0", ROLE_LABELS[ROLE_NOBET_BASI])]
    if is_day_shift:
        columns.append((f"{ROLE_KAPICI}#0", ROLE_LABELS[ROLE_KAPICI]))

    if max_comez_basi == 1:
        columns.append((f"{ROLE_COMEZ_BASI}#0", ROLE_LABELS[ROLE_COMEZ_BASI]))
    else:
        for i in range(max_comez_basi):
            columns.append((f"{ROLE_COMEZ_BASI}#{i}", f"{ROLE_LABELS[ROLE_COMEZ_BASI]} {i + 1}"))

    for i in range(max_comez):
        columns.append((f"{ROLE_COMEZ}#{i}", f"Çömez {i + 1}"))

    return columns


def _cell_value(result: ScheduleResult, shift: Shift, column_key: str) -> str:
    base_role, index = column_key.split("#")
    names = result.assignments.get((shift, base_role), [])
    index = int(index)
    return names[index] if index < len(names) else ""


def _write_schedule_sheet(ws: Worksheet, result: ScheduleResult) -> None:
    ws.title = "Nöbet Listesi"

    day_columns = _shift_columns(result, is_day_shift=True)
    night_columns = _shift_columns(result, is_day_shift=False)

    ws.cell(row=1, column=1, value="Tarih")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    col = 2
    ws.cell(row=1, column=col, value="GÜNDÜZ")
    ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(day_columns) - 1)
    for _, label in day_columns:
        ws.cell(row=2, column=col, value=label)
        col += 1

    ws.cell(row=1, column=col, value="GECE")
    ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(night_columns) - 1)
    for _, label in night_columns:
        ws.cell(row=2, column=col, value=label)
        col += 1

    last_col = col
    for c in range(1, last_col):
        for r in (1, 2):
            cell = ws.cell(row=r, column=c)
            cell.alignment = CENTER
            if r == 1:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
            else:
                cell.font = GROUP_HEADER_FONT

    day_shifts_by_date = {s.the_date: s for s in result.shifts if s.is_day_shift}
    night_shifts_by_date = {s.the_date: s for s in result.shifts if not s.is_day_shift}

    row = 3
    for the_date in sorted(day_shifts_by_date):
        ws.cell(row=row, column=1, value=the_date.isoformat())
        c = 2
        day_shift = day_shifts_by_date[the_date]
        for column_key, _ in day_columns:
            ws.cell(row=row, column=c, value=_cell_value(result, day_shift, column_key))
            c += 1
        night_shift = night_shifts_by_date[the_date]
        for column_key, _ in night_columns:
            ws.cell(row=row, column=c, value=_cell_value(result, night_shift, column_key))
            c += 1
        row += 1

    for c in range(1, last_col):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.freeze_panes = "B3"

    _write_footer(ws, last_used_row=row - 1)


def _day_night_counts(result: ScheduleResult) -> dict[str, tuple[int, int]]:
    """Per-doctor (day_count, night_count), tallied from every role assignment
    across every shift — day_count + night_count always equals doctor_totals."""
    counts: dict[str, list[int]] = {}
    for (shift, _role), names in result.assignments.items():
        index = 0 if shift.is_day_shift else 1
        for name in names:
            counts.setdefault(name, [0, 0])[index] += 1
    return {name: (day, night) for name, (day, night) in counts.items()}


def _write_summary_sheet(ws: Worksheet, result: ScheduleResult, doctors: list[Doctor]) -> None:
    ws.title = "Özet"
    headers = [
        "Doktor",
        "Kıdem",
        "Hedef Nöbet Sayısı",
        "Atanan Nöbet Sayısı",
        "Gündüz Nöbet Sayısı",
        "Gece Nöbet Sayısı",
    ]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER

    day_night_counts = _day_night_counts(result)

    for r, d in enumerate(doctors, start=2):
        assigned = result.doctor_totals.get(d.name, 0)
        day_count, night_count = day_night_counts.get(d.name, (0, 0))
        ws.cell(row=r, column=1, value=d.name)
        ws.cell(row=r, column=2, value=d.seniority)
        ws.cell(row=r, column=3, value=d.shift_count_target)
        ws.cell(row=r, column=4, value=assigned)
        ws.cell(row=r, column=5, value=day_count)
        ws.cell(row=r, column=6, value=night_count)

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 22

    _write_footer(ws, last_used_row=len(doctors) + 1)
